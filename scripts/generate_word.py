from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.shared import Pt
import re

# 创建文档
doc = Document()

# 设置页面大小为A4
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)

# 读取Markdown内容
with open('操作说明书.md', 'r', encoding='utf-8') as f:
    content = f.read()

# 解析Markdown并转换为Word
lines = content.split('\n')

for line in lines:
    line = line.strip()
    if not line:
        continue
    
    # 处理标题
    if line.startswith('#'):
        level = line.count('#')
        text = line.strip('#').strip()
        if level == 1:
            heading = doc.add_heading(text, level=0)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            doc.add_heading(text, level=level-1)
    
    # 处理列表项
    elif line.startswith('* ') or line.startswith('- '):
        text = line[2:]
        doc.add_paragraph(text, style='ListBullet')
    
    elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. ') or line.startswith('4. ') or line.startswith('5. '):
        text = line[3:]
        doc.add_paragraph(text, style='ListNumber')
    
    # 处理普通段落
    else:
        doc.add_paragraph(line)

# 保存文档
doc.save('操作说明书.docx')
print('Word文档生成完成')